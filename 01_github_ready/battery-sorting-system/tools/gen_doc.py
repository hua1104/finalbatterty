# -*- coding: utf-8 -*-
from docx import Document

out_path = '锂电池外观缺陷智能检测与分拣系统 用户手册 V1.0.docx'

_doc = Document()
_doc.add_paragraph('锂电池外观缺陷智能检测与分拣系统')
_doc.add_paragraph('V1.0')
_doc.add_paragraph('用户手册')

_doc.add_paragraph('')
_doc.add_paragraph('目录')

_doc.add_paragraph('第一章\t系统简介')
_doc.add_paragraph('1.\t背景')
_doc.add_paragraph('2.\t简介')
_doc.add_paragraph('3.\t目标用户')
_doc.add_paragraph('4.\t主要功能')
_doc.add_paragraph('5.\t优势与创新')

_doc.add_paragraph('第二章\t系统运行环境')
_doc.add_paragraph('1.\t系统硬件环境')
_doc.add_paragraph('2.\t系统软件环境')

_doc.add_paragraph('第三章\t系统主要功能（用户）')
_doc.add_paragraph('1.\t登录与权限验证')
_doc.add_paragraph('2.\t实时监控与AI判定')
_doc.add_paragraph('3.\t生产控制与自动分拣')
_doc.add_paragraph('4.\t数据看板与统计')
_doc.add_paragraph('5.\t追溯记录与历史查看')
_doc.add_paragraph('6.\t设备状态与健康监测')

_doc.add_paragraph('第四章\t系统主要功能（管理员）')
_doc.add_paragraph('1.\t用户与角色管理')
_doc.add_paragraph('2.\t记录查询与筛选')
_doc.add_paragraph('3.\t复核修正与结果确认')
_doc.add_paragraph('4.\t数据导出与报表')
_doc.add_paragraph('5.\t数据清理与批量删除')

_doc.add_heading('系统简介', level=1)
_doc.add_heading('背景', level=2)
_doc.add_paragraph('锂电池外观质量直接影响安全性与一致性。传统人工目检存在效率低、主观性强、可追溯性差的问题，难以满足规模化生产线的稳定检测需求。为提升检测效率、降低误判率并实现可追溯管理，需要引入基于视觉与自动化分拣的智能检测系统。')

_doc.add_heading('简介', level=2)
_doc.add_paragraph('本系统采用“边缘端 + 云端”架构。边缘端部署在生产线（Raspberry Pi），负责摄像头采集、AI缺陷识别、结果显示与自动分拣；云端部署在服务器（Linux），负责数据接收、统计分析、历史追溯与用户权限管理。系统集成了缺陷检测、实时监控、数据看板与历史管理等能力，实现了生产线外观质检流程的自动化与数据化。')

_doc.add_heading('目标用户', level=2)
_doc.add_paragraph('主要面向锂电池生产线操作员、质量工程师、工艺与设备维护人员以及生产管理人员。')

_doc.add_heading('主要功能', level=2)
_doc.add_paragraph('实时采集与AI缺陷识别')
_doc.add_paragraph('自动分拣与生产启停控制')
_doc.add_paragraph('缺陷统计与良率分析')
_doc.add_paragraph('历史记录追溯与图片查看')
_doc.add_paragraph('云端数据管理与权限控制')

_doc.add_heading('优势与创新', level=2)
_doc.add_paragraph('边缘侧低时延判定，满足实时生产节拍')
_doc.add_paragraph('AI识别与机械执行联动，实现闭环分拣')
_doc.add_paragraph('云端集中存储与统计分析，支持跨班次追溯')
_doc.add_paragraph('本地/云端双重登录校验，保障现场连续运行')
_doc.add_paragraph('可扩展缺陷类别与统计维度，便于后续迭代')

_doc.add_heading('系统运行环境', level=1)
_doc.add_heading('系统硬件环境', level=2)
_doc.add_paragraph('边缘端：Raspberry Pi 4B/5、USB摄像头、红外传感器、步进电机与驱动器（传送带与分拣机构）、生产线传送带、稳定电源。')
_doc.add_paragraph('云端：Linux服务器（Ubuntu/CentOS），用于部署Web服务与MySQL数据库。')

_doc.add_heading('系统软件环境', level=2)
_doc.add_paragraph('边缘端系统：Raspberry Pi OS，Python 3.8+。')
_doc.add_paragraph('云端系统：Ubuntu/CentOS，Python 3.8+，MySQL 5.7/8.0。')
_doc.add_paragraph('主要依赖：OpenCV、Ultralytics YOLO、Pillow、Requests、Flask、Waitress、Pandas、OpenPyXL、ECharts（前端图表）。')

_doc.add_heading('系统主要功能（用户）', level=1)
_doc.add_heading('登录与权限验证', level=2)
_doc.add_paragraph('系统支持本地与云端双重校验。现场断网情况下可使用本地账号登录，联网时自动进行云端校验。')

_doc.add_heading('实时监控与AI判定', level=2)
_doc.add_paragraph('实时显示摄像头画面与AI判定结果，缺陷目标以标注框方式可视化，便于操作员确认。')

_doc.add_heading('生产控制与自动分拣', level=2)
_doc.add_paragraph('支持一键启动/停止生产线。检测为不合格时自动触发分拣机构剔除，合格品继续通过。')

_doc.add_heading('数据看板与统计', level=2)
_doc.add_paragraph('提供合格/不合格数量统计与良率分析，实时更新生产数据，辅助现场质量判断。')

_doc.add_heading('追溯记录与历史查看', level=2)
_doc.add_paragraph('展示最近检测记录，包括时间、结果、缺陷类型与置信度，便于快速追溯与复查。')

_doc.add_heading('设备状态与健康监测', level=2)
_doc.add_paragraph('展示云端连接、传感器、相机、电机与系统负载等状态，帮助定位设备异常。')

_doc.add_heading('系统主要功能（管理员）', level=1)
_doc.add_heading('用户与角色管理', level=2)
_doc.add_paragraph('管理员可新增、修改、删除用户，并分配操作员/管理员角色，统一管理登录权限。')

_doc.add_heading('记录查询与筛选', level=2)
_doc.add_paragraph('支持按日期范围、操作员与结果类型查询检测记录，便于统计与追溯。')

_doc.add_heading('复核修正与结果确认', level=2)
_doc.add_paragraph('支持人工复核并修正AI判定结果，标记复核状态，保证统计准确性。')

_doc.add_heading('数据导出与报表', level=2)
_doc.add_paragraph('一键导出检测记录与统计结果，生成Excel报表，便于质量管理留档。')

_doc.add_heading('数据清理与批量删除', level=2)
_doc.add_paragraph('支持按时间范围批量删除历史记录及图片文件，便于存储管理。')

_doc.save(out_path)
print(out_path)
